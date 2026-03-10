from fastapi.responses import StreamingResponse
import io

def export_to_excel(resume: dict, jobs: list):
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return {"error": "openpyxl not installed"}

    wb = openpyxl.Workbook()

    # Sheet 1: Jobs
    ws = wb.active
    ws.title = "Matched Jobs"
    headers = ["#", "Title", "Company", "Location", "Match %", "Matched Skills", "Portal", "Apply Link", "Salary", "Posted"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(horizontal="center")

    for i, job in enumerate(jobs, 1):
        ws.append([
            i,
            job.get("title", ""),
            job.get("company", ""),
            job.get("location", ""),
            f"{job.get('match_score', 0)}%",
            ", ".join(job.get("matched_skills", [])),
            job.get("portal", ""),
            job.get("apply_link", ""),
            job.get("salary", ""),
            job.get("posted_date", ""),
        ])

    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 20

    # Sheet 2: Resume Skills
    ws2 = wb.create_sheet("My Skills")
    ws2.append(["Skill"])
    for sk in resume.get("skills", []):
        ws2.append([sk])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=jobradar_results.xlsx"}
    )
